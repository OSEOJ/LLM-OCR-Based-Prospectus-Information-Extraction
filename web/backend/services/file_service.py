import os
import json
import uuid
from datetime import datetime
from typing import Dict, Any, Optional
from pathlib import Path
import fitz  # PyMuPDF
from docx import Document
import logging

logger = logging.getLogger(__name__)

class FileService:
    def __init__(self):
        self.output_dir = os.getenv("OUTPUT_DIR", "./output")
        self.temp_dir = os.getenv("TEMP_DIR", "./temp")
        
        # 디렉터리 생성
        os.makedirs(self.output_dir, exist_ok=True)
        os.makedirs(self.temp_dir, exist_ok=True)
    
    def is_ready(self) -> bool:
        """파일 서비스 준비 상태 확인"""
        return (os.path.exists(self.output_dir) and 
                os.path.exists(self.temp_dir) and
                os.access(self.output_dir, os.W_OK) and
                os.access(self.temp_dir, os.W_OK))
    
    def generate_file_id(self) -> str:
        """고유한 파일 ID 생성"""
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        unique_id = str(uuid.uuid4())[:8]
        return f"{timestamp}_{unique_id}"
    
    def save_conversion_result(self, 
                             original_filename: str,
                             txt_content: str,
                             json_result: Dict[str, Any],
                             product_type: str) -> str:
        """변환 결과를 파일로 저장하고 파일 ID 반환"""
        file_id = self.generate_file_id()
        
        # 메타데이터 생성
        metadata = {
            "file_id": file_id,
            "original_filename": original_filename,
            "product_type": product_type,
            "created_at": datetime.now().isoformat(),
            "txt_length": len(txt_content),
            "json_fields_count": len(json_result) if isinstance(json_result, dict) else 0
        }
        
        # 파일 저장
        txt_path = os.path.join(self.output_dir, f"{file_id}.txt")
        with open(txt_path, 'w', encoding='utf-8') as f:
            f.write(txt_content)
        
        json_path = os.path.join(self.output_dir, f"{file_id}.json")
        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(json_result, f, ensure_ascii=False, indent=4)
        
        metadata_path = os.path.join(self.output_dir, f"{file_id}_meta.json")
        with open(metadata_path, 'w', encoding='utf-8') as f:
            json.dump(metadata, f, ensure_ascii=False, indent=4)
        
        return file_id
    
    def get_result_file_path(self, file_id: str, format: str) -> str:
        """결과 파일 경로 반환"""
        if format not in ['json', 'txt', 'meta']:
            raise ValueError("지원되는 형식: json, txt, meta")
        
        if format == 'meta':
            filename = f"{file_id}_meta.json"
        else:
            filename = f"{file_id}.{format}"
        
        file_path = os.path.join(self.output_dir, filename)
        
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"파일을 찾을 수 없습니다: {filename}")
        
        return file_path
    
    def get_conversion_metadata(self, file_id: str) -> Dict[str, Any]:
        """변환 메타데이터 반환"""
        metadata_path = self.get_result_file_path(file_id, 'meta')
        
        with open(metadata_path, 'r', encoding='utf-8') as f:
            metadata = json.load(f)
        
        return metadata
    
    def list_conversion_results(self, limit: int = 50) -> list:
        """최근 변환 결과 목록 반환"""
        # 메타데이터 파일들 찾기
        meta_files = []
        for filename in os.listdir(self.output_dir):
            if filename.endswith('_meta.json'):
                file_path = os.path.join(self.output_dir, filename)
                if os.path.isfile(file_path):
                    meta_files.append(file_path)
        
        # 생성 시간 기준으로 정렬
        meta_files.sort(key=lambda x: os.path.getctime(x), reverse=True)
        
        # 메타데이터 로드
        results = []
        for meta_file in meta_files[:limit]:
            try:
                with open(meta_file, 'r', encoding='utf-8') as f:
                    metadata = json.load(f)
                results.append(metadata)
            except Exception:
                continue
        
        return results
    
    def delete_conversion_result(self, file_id: str) -> bool:
        """변환 결과 파일들 삭제"""
        deleted_files = []
        file_extensions = ['txt', 'json', '_meta.json']
        
        for ext in file_extensions:
            if ext == '_meta.json':
                filename = f"{file_id}_meta.json"
            else:
                filename = f"{file_id}.{ext}"
            
            file_path = os.path.join(self.output_dir, filename)
            
            if os.path.exists(file_path):
                os.remove(file_path)
                deleted_files.append(filename)
        
        return len(deleted_files) > 0
    
    def cleanup_old_files(self, days: int = 7) -> int:
        """오래된 파일들 정리 (기본 7일)"""
        from datetime import timedelta
        
        cutoff_time = datetime.now() - timedelta(days=days)
        deleted_count = 0
        
        for filename in os.listdir(self.output_dir):
            file_path = os.path.join(self.output_dir, filename)
            
            if os.path.isfile(file_path):
                file_time = datetime.fromtimestamp(os.path.getctime(file_path))
                
                if file_time < cutoff_time:
                    try:
                        os.remove(file_path)
                        deleted_count += 1
                    except Exception:
                        continue
        
        return deleted_count
    
    def get_storage_info(self) -> Dict[str, Any]:
        """저장소 정보 반환"""
        total_files = 0
        total_size = 0
        
        for filename in os.listdir(self.output_dir):
            file_path = os.path.join(self.output_dir, filename)
            if os.path.isfile(file_path):
                total_files += 1
                total_size += os.path.getsize(file_path)
        
        return {
            "output_directory": self.output_dir,
            "temp_directory": self.temp_dir,
            "total_files": total_files,
            "total_size_bytes": total_size,
            "total_size_mb": round(total_size / (1024 * 1024), 2)
        }
    
    def needs_ocr(self, pdf_path: str) -> bool:
        """PDF가 OCR이 필요한지 확인"""
        try:
            doc = fitz.open(pdf_path)
            for page in doc:
                text = page.get_text().strip()
                if len(text) > 50:  # 50자 이상의 텍스트가 있으면 텍스트 추출 가능
                    doc.close()
                    return False
            doc.close()
            return True
        except Exception as e:
            logger.error(f"[OCR 판단 오류] {pdf_path} → {e}")
            return True
    
    def extract_pdf_text(self, pdf_path: str) -> str:
        """PDF에서 텍스트 직접 추출"""
        try:
            validation = self.is_valid_pdf(pdf_path)
            if not validation["is_valid"]:
                logger.error(f"PDF 파일 유효성 검사 실패: {validation['error']}")
                return ""
            
            doc = fitz.open(pdf_path)
            full_text = ""
            for page_num in range(len(doc)):
                page = doc[page_num]
                page_text = page.get_text("text")
                if page_text.strip():
                    full_text += f"[페이지 {page_num + 1}]\n{page_text}\n\n"
            
            doc.close()
            extracted_text = full_text.strip()
            
            logger.info(f"PDF 텍스트 추출 완료: {os.path.basename(pdf_path)} - 페이지: {validation['page_count']}, 텍스트 길이: {len(extracted_text)}")
            return extracted_text
            
        except Exception as e:
            logger.error(f"[PDF 오류] {pdf_path}: {e}")
            return ""
    
    def extract_docx_text(self, docx_path: str) -> str:
        """DOCX에서 텍스트 추출"""
        try:
            validation = self.is_valid_docx(docx_path)
            if not validation["is_valid"]:
                logger.error(f"DOCX 파일 유효성 검사 실패: {validation['error']}")
                return ""
            
            if not validation["has_content"]:
                logger.warning(f"DOCX 파일에 텍스트 내용이 없습니다: {docx_path}")
                return ""
            
            doc = Document(docx_path)
            full_text = []

            # 일반 문단 텍스트 추출
            paragraph_count = 0
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
                    paragraph_count += 1

            # 표 텍스트 추출
            table_count = 0
            for table in doc.tables:
                table_count += 1
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        full_text.append(" | ".join(row_text))

            extracted_text = "\n".join(full_text).strip()
            logger.info(f"DOCX 텍스트 추출 완료: {os.path.basename(docx_path)} - 문단: {paragraph_count}, 표: {table_count}, 텍스트 길이: {len(extracted_text)}")
            
            return extracted_text
            
        except Exception as e:
            logger.error(f"[DOCX 오류] {docx_path}: {e}")
            return ""
    
    def determine_processing_method(self, file_path: str) -> str:
        """파일 유형에 따라 처리 방법 결정"""
        file_ext = Path(file_path).suffix.lower()
        
        if file_ext == ".pdf":
            # PDF인 경우 텍스트 추출 가능성 확인
            if self.needs_ocr(file_path):
                return "ocr"
            else:
                return "text_extraction"
        elif file_ext == ".docx":
            # DOCX는 항상 텍스트 추출
            return "text_extraction"
        else:
            # 기타 파일은 OCR
            return "ocr"
    
    def is_valid_docx(self, docx_path: str) -> Dict[str, Any]:
        """DOCX 파일 유효성 검사"""
        result = {
            "is_valid": False,
            "error": None,
            "file_size": 0,
            "has_content": False
        }
        
        if not os.path.exists(docx_path):
            result["error"] = "파일이 존재하지 않습니다"
            return result
        
        file_size = os.path.getsize(docx_path)
        result["file_size"] = file_size
        
        if file_size < 1000:  # 1KB 미만
            result["error"] = "파일이 너무 작습니다"
            return result
        
        try:
            doc = Document(docx_path)
            
            # 텍스트 내용 확인
            text_content = ""
            for para in doc.paragraphs:
                if para.text.strip():
                    text_content += para.text + "\n"
            
            # 표 내용도 확인
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_content += cell.text + "\n"
            
            result["has_content"] = len(text_content.strip()) > 0
            result["is_valid"] = True
            
        except Exception as e:
            result["error"] = f"DOCX 파일 검증 오류: {str(e)}"
        
        return result
    
    def is_valid_pdf(self, pdf_path: str) -> Dict[str, Any]:
        """PDF 파일 유효성 검사"""
        result = {
            "is_valid": False,
            "error": None,
            "file_size": 0,
            "page_count": 0
        }
        
        if not os.path.exists(pdf_path):
            result["error"] = "파일이 존재하지 않습니다"
            return result
        
        file_size = os.path.getsize(pdf_path)
        result["file_size"] = file_size
        
        if file_size < 1000:  # 1KB 미만
            result["error"] = "파일이 너무 작습니다"
            return result
        
        try:
            doc = fitz.open(pdf_path)
            result["page_count"] = len(doc)
            doc.close()
            
            if result["page_count"] == 0:
                result["error"] = "PDF에 페이지가 없습니다"
                return result
            
            result["is_valid"] = True
            
        except Exception as e:
            result["error"] = f"PDF 파일 검증 오류: {str(e)}"
        
        return result